import io
from unittest.mock import patch, Mock
import pymupdf
import docx
import requests

from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from rest_framework import status
from rest_framework.test import APITestCase, APITransactionTestCase

from .models import Document

def create_sample_pdf_bytes(content: str = "This is a sample MRPL report text.") -> bytes:
    """Generates valid minimal in-memory PDF bytes with text."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 72), content)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes

class DocumentAPITests(APITestCase):
    def setUp(self):
        self.upload_url = reverse("document-upload")
        self.list_url = reverse("document-list")

    def test_upload_valid_txt_file(self):
        """Test uploading a valid TXT file extracts text and creates Document."""
        txt_content = "MRPL Confidential Guidelines: All data stays on PC5."
        txt_file = SimpleUploadedFile("guidelines.txt", txt_content.encode("utf-8"), content_type="text/plain")

        response = self.client.post(self.upload_url, {"file": txt_file}, format="multipart")

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertTrue(response.data["success"])
        self.assertEqual(response.data["document"]["original_name"], "guidelines.txt")
        self.assertEqual(response.data["document"]["file_type"], "TXT")
        self.assertIn("MRPL Confidential Guidelines", response.data["document"]["extracted_text"])

    def test_upload_valid_pdf_file(self):
        """Test uploading a valid PDF extracts text via PyMuPDF."""
        pdf_bytes = create_sample_pdf_bytes("Phase 2 local document assistant operational verification.")
        pdf_file = SimpleUploadedFile("report.pdf", pdf_bytes, content_type="application/pdf")

        response = self.client.post(self.upload_url, {"file": pdf_file}, format="multipart")

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertTrue(response.data["success"])
        self.assertEqual(response.data["document"]["file_type"], "PDF")
        self.assertIn("Phase 2 local document assistant", response.data["document"]["extracted_text"])

    def test_reject_unsupported_extensions(self):
        """Test that extensions like .docx, .xlsx, .gif, .bmp are rejected (PNG/JPG now go to OCR)."""
        bad_files = [
            ("data.docx", b"dummy docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            ("sheet.xlsx", b"dummy xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            ("anim.gif", b"GIF89a", "image/gif"),
            ("scan.bmp", b"BM", "image/bmp"),
        ]

        for filename, content, mime in bad_files:
            file_obj = SimpleUploadedFile(filename, content, content_type=mime)
            response = self.client.post(self.upload_url, {"file": file_obj}, format="multipart")
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
            self.assertFalse(response.data.get("success", True))
            self.assertIn("Unsupported file type", response.data.get("error", ""))

    def test_reject_empty_file(self):
        """Test that empty TXT or PDF files are rejected with 400."""
        empty_txt = SimpleUploadedFile("empty.txt", b"", content_type="text/plain")
        response = self.client.post(self.upload_url, {"file": empty_txt}, format="multipart")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("empty", response.data.get("error", "").lower())

    def test_character_limit_truncation(self):
        """Test that documents exceeding 50,000 characters are capped."""
        long_content = "A" * 60000
        txt_file = SimpleUploadedFile("huge.txt", long_content.encode("utf-8"), content_type="text/plain")

        response = self.client.post(self.upload_url, {"file": txt_file}, format="multipart")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        extracted = response.data["document"]["extracted_text"]
        self.assertTrue(extracted.startswith("A" * 50000))
        self.assertIn("truncated at 50,000 characters", extracted)

    @patch("chat.model_manager.requests.Session.post")
    def test_ask_question_about_document(self, mock_post):
        """Test asking a question about an uploaded document."""
        mock_resp = Mock()
        mock_resp.status_code = 200
        mock_resp.json.return_value = {"response": "The document specifies on-premise inference on PC5.", "message": {"role": "assistant", "content": "The document specifies on-premise inference on PC5."}}
        mock_resp.raise_for_status = Mock()
        mock_post.return_value = mock_resp

        doc = Document.objects.create(
            original_name="test.txt",
            file_type="TXT",
            extracted_text="MRPL system uses PC5 workstation with local Qwen model.",
        )

        ask_url = reverse("document-ask", kwargs={"pk": doc.id})
        response = self.client.post(ask_url, {"question": "Where does the model run?"}, format="json")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertTrue(response.data["success"])
        self.assertEqual(response.data["answer"], "The document specifies on-premise inference on PC5.")

    @patch("chat.model_manager.requests.Session.post")
    def test_summarize_document(self, mock_post):
        """Test generating an AI summary of an uploaded document."""
        mock_resp = Mock()
        mock_resp.status_code = 200
        mock_resp.json.return_value = {"response": "1. Executive Overview: Project completed successfully.", "message": {"role": "assistant", "content": "1. Executive Overview: Project completed successfully."}}
        mock_resp.raise_for_status = Mock()
        mock_post.return_value = mock_resp

        doc = Document.objects.create(
            original_name="summary_test.txt",
            file_type="TXT",
            extracted_text="Project completed on schedule with zero cloud dependencies.",
        )

        sum_url = reverse("document-summary", kwargs={"pk": doc.id})
        response = self.client.post(sum_url, {}, format="json")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertTrue(response.data["success"])
        self.assertIn("Executive Overview", response.data["summary"])

    @patch("chat.model_manager.requests.Session.post")
    def test_ollama_offline_returns_503(self, mock_post):
        """Test that Ollama connection error returns HTTP 503."""
        mock_post.side_effect = requests.exceptions.ConnectionError("Connection refused")

        doc = Document.objects.create(
            original_name="offline.txt",
            file_type="TXT",
            extracted_text="Sample text.",
        )

        ask_url = reverse("document-ask", kwargs={"pk": doc.id})
        response = self.client.post(ask_url, {"question": "Hello?"}, format="json")

        self.assertEqual(response.status_code, status.HTTP_503_SERVICE_UNAVAILABLE)
        self.assertFalse(response.data["success"])
        self.assertIn("unavailable", response.data["error"].lower())

    def test_generate_and_download_docx_report(self):
        """Test generating a Word DOCX report returns valid docx binary."""
        doc = Document.objects.create(
            original_name="audit_report.txt",
            file_type="TXT",
            extracted_text="Detailed audit findings for MRPL infrastructure.",
        )

        report_url = reverse("document-generate-report", kwargs={"pk": doc.id})
        response = self.client.post(
            report_url,
            {"summary": "Summary: All systems meet on-premise compliance standards."},
            format="json"
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        self.assertIn("attachment; filename=", response["Content-Disposition"])

        # Parse generated docx bytes to verify valid Word document structure
        doc_stream = io.BytesIO(response.content)
        parsed_docx = docx.Document(doc_stream)
        text_content = "\n".join([p.text for p in parsed_docx.paragraphs])
        self.assertIn("MRPL AI Workbench", text_content)
        self.assertIn("audit_report.txt", [cell.text for row in parsed_docx.tables[0].rows for cell in row.cells])

    def test_delete_document(self):
        """Test deleting a document removes it from the database."""
        doc = Document.objects.create(
            original_name="delete_me.txt",
            file_type="TXT",
            extracted_text="Temp text",
        )

        detail_url = reverse("document-detail", kwargs={"pk": doc.id})
        response = self.client.delete(detail_url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertFalse(Document.objects.filter(pk=doc.id).exists())

    def test_phase1_chat_still_works(self):
        """Regression test: Phase 1 /api/chat/ endpoint still works."""
        with patch("chat.model_manager.requests.Session.post") as mock_post:
            mock_resp = Mock()
            mock_resp.status_code = 200
            mock_resp.json.return_value = {"response": "Phase 1 chat response", "message": {"role": "assistant", "content": "Phase 1 chat response"}}
            mock_resp.raise_for_status = Mock()
            mock_post.return_value = mock_resp

            response = self.client.post(
                reverse("chat"),
                {"message": "Hello chat"},
                format="json"
            )
            self.assertEqual(response.status_code, status.HTTP_200_OK)
            self.assertTrue(response.data["success"])


# ---------------------------------------------------------------------------
# Step 10: local OCR + routed document reasoning (PaddleOCR mocked for speed)
# ---------------------------------------------------------------------------
import base64
import shutil
import tempfile

from django.test import override_settings

TEMP_MEDIA = tempfile.mkdtemp(prefix="mrpl-test-media-")


def make_png_bytes(width=400, height=200) -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, width, height), False)
    pix.clear_with(255)
    return pix.tobytes("png")


def make_scanned_pdf_bytes(pages=1, text_pages=()) -> bytes:
    """PDF whose pages are images only (no text layer), except indexes in text_pages."""
    doc = pymupdf.open()
    for i in range(pages):
        page = doc.new_page()
        if i in text_pages:
            page.insert_text((50, 72), f"Typed page {i + 1}: pressure test passed at 12 bar.")
        else:
            page.insert_image(page.rect, stream=make_png_bytes(300, 400))
    data = doc.tobytes()
    doc.close()
    return data


def chat_prompt(payload):
    """All text sent to Ollama (/api/chat messages or /api/generate prompt)."""
    if "messages" in payload:
        return "\n".join(m["content"] for m in payload["messages"])
    return payload["prompt"]


def chat_images(payload):
    if "messages" in payload:
        return payload["messages"][-1].get("images")
    return payload.get("images")


def ok_response(text):
    resp = Mock()
    resp.status_code = 200
    resp.json.return_value = {"response": text, "message": {"role": "assistant", "content": text}}
    resp.raise_for_status = Mock()
    return resp


OCR_RESULT = {"text": "Inspection Date: 14 Sep 2026\nFinding: shell corrosion", "lines": 2, "confidence": 0.98}


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class DocumentOCRTests(APITestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(TEMP_MEDIA, ignore_errors=True)

    def upload(self, name, data, mime):
        return self.client.post(
            reverse("document-upload"),
            {"file": SimpleUploadedFile(name, data, content_type=mime)},
            format="multipart",
        )

    @patch("documents.ocr.ocr_image", return_value=OCR_RESULT)
    def test_png_upload_runs_local_ocr(self, mock_ocr):
        response = self.upload("scan.png", make_png_bytes(), "image/png")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        doc = response.data["document"]
        self.assertEqual(doc["file_type"], "PNG")
        self.assertEqual(doc["extraction_method"], "ocr")
        self.assertEqual(doc["ocr_pages"], 1)
        self.assertEqual(doc["ocr_confidence"], 0.98)
        self.assertIn("shell corrosion", doc["extracted_text"])
        mock_ocr.assert_called_once()

    @patch("documents.ocr.ocr_image", return_value=OCR_RESULT)
    def test_jpeg_extension_accepted(self, mock_ocr):
        response = self.upload("scan.jpeg", b"\xff\xd8\xff\xe0" + b"\x00" * 20, "image/jpeg")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["document"]["file_type"], "JPG")

    @patch("documents.ocr.ocr_image")
    def test_fake_image_rejected_before_ocr(self, mock_ocr):
        response = self.upload("fake.png", b"not really a png", "image/png")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        mock_ocr.assert_not_called()

    @patch("documents.ocr.ocr_pdf_page", return_value=OCR_RESULT)
    def test_scanned_pdf_pages_are_ocred(self, mock_ocr):
        response = self.upload("scan.pdf", make_scanned_pdf_bytes(pages=2), "application/pdf")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        doc = response.data["document"]
        self.assertEqual(doc["extraction_method"], "ocr")
        self.assertEqual(doc["ocr_pages"], 2)
        self.assertIn("--- Page 1 (OCR) ---", doc["extracted_text"])
        self.assertIn("--- Page 2 (OCR) ---", doc["extracted_text"])
        self.assertEqual(mock_ocr.call_count, 2)

    @patch("documents.ocr.ocr_pdf_page", return_value=OCR_RESULT)
    def test_mixed_pdf_only_ocrs_scanned_pages(self, mock_ocr):
        response = self.upload("mixed.pdf", make_scanned_pdf_bytes(pages=2, text_pages=(0,)), "application/pdf")
        doc = response.data["document"]
        self.assertEqual(doc["extraction_method"], "text+ocr")
        self.assertEqual(doc["ocr_pages"], 1)
        self.assertIn("pressure test passed", doc["extracted_text"])
        self.assertIn("--- Page 2 (OCR) ---", doc["extracted_text"])
        mock_ocr.assert_called_once()

    @patch("documents.ocr.ocr_pdf_page")
    def test_text_pdf_does_not_use_ocr(self, mock_ocr):
        response = self.upload("typed.pdf", create_sample_pdf_bytes("A normal digital PDF with a text layer."), "application/pdf")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["document"]["extraction_method"], "text")
        mock_ocr.assert_not_called()

    @override_settings(OCR_MAX_PDF_PAGES=1)
    @patch("documents.ocr.ocr_pdf_page", return_value=OCR_RESULT)
    def test_ocr_page_cap(self, mock_ocr):
        response = self.upload("long.pdf", make_scanned_pdf_bytes(pages=3), "application/pdf")
        self.assertEqual(response.data["document"]["ocr_pages"], 1)
        self.assertIn("2 further scanned page(s) were not processed", response.data["document"]["extracted_text"])

    @patch("documents.ocr.ocr_pdf_page", return_value={"text": "", "lines": 0, "confidence": 0.0})
    def test_blank_scan_returns_clean_error(self, mock_ocr):
        response = self.upload("blank.pdf", make_scanned_pdf_bytes(pages=1), "application/pdf")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("even with local OCR", response.data["error"])

    @patch("chat.model_manager.requests.Session.post", return_value=ok_response("Key findings: corrosion."))
    def test_ocr_text_question_routes_to_qwen3(self, mock_post):
        doc = Document.objects.create(
            original_name="scan.pdf", file_type="PDF", extraction_method="ocr",
            ocr_pages=1, extracted_text="Finding: shell corrosion near saddle.",
        )
        response = self.client.post(
            reverse("document-ask", kwargs={"pk": doc.id}),
            {"question": "Summarize this document and list the key findings."}, format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["ollama_model"], "qwen3:8b")
        self.assertEqual({k: v for k, v in response.data["routing"].items() if k != "explanation"}, {"mode": "auto", "reason": "general_reasoning"})
        payload = mock_post.call_args.kwargs["json"]
        self.assertIn("shell corrosion near saddle", chat_prompt(payload))
        self.assertIn("produced by local OCR", chat_prompt(payload))
        self.assertIsNone(chat_images(payload))

    @patch("chat.model_manager.requests.Session.post", return_value=ok_response("A pump symbol."))
    def test_visual_question_sends_image_to_vision_model(self, mock_post):
        doc = Document.objects.create(
            original_name="drawing.png", file_type="PNG", extraction_method="ocr",
            extracted_text="P-101", file=SimpleUploadedFile("drawing.png", make_png_bytes(2400, 1200)),
        )
        response = self.client.post(
            reverse("document-ask", kwargs={"pk": doc.id}),
            {"question": "Describe the diagram in this image."}, format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["ollama_model"], "qwen2.5vl:7b")
        self.assertEqual(response.data["routing"]["reason"], "vision")
        payload = mock_post.call_args.kwargs["json"]
        self.assertEqual(len(chat_images(payload)), 1)
        sent = pymupdf.Pixmap(base64.b64decode(chat_images(payload)[0]))
        self.assertLessEqual(max(sent.width, sent.height), 1024)  # downscaled for 8 GB VRAM

    @patch("chat.model_manager.requests.Session.post", return_value=ok_response("A title block."))
    def test_visual_question_on_scanned_pdf_sends_page_one(self, mock_post):
        doc = Document.objects.create(
            original_name="scan.pdf", file_type="PDF", extraction_method="ocr", extracted_text="x" * 30,
            file=SimpleUploadedFile("scan.pdf", make_scanned_pdf_bytes(pages=2)),
        )
        response = self.client.post(
            reverse("document-ask", kwargs={"pk": doc.id}),
            {"question": "What does the drawing show?"}, format="json",
        )
        self.assertEqual(response.data["ollama_model"], "qwen2.5vl:7b")
        self.assertEqual(len(chat_images(mock_post.call_args.kwargs["json"])), 1)

    @patch("chat.model_manager.requests.Session.post", return_value=ok_response("ok"))
    def test_visual_question_on_txt_falls_back_to_text_model(self, mock_post):
        doc = Document.objects.create(original_name="n.txt", file_type="TXT", extracted_text="notes")
        response = self.client.post(
            reverse("document-ask", kwargs={"pk": doc.id}),
            {"question": "Describe the image."}, format="json",
        )
        self.assertEqual(response.data["ollama_model"], "qwen3:8b")
        self.assertEqual(response.data["routing"]["reason"], "general_reasoning")
        self.assertIsNone(chat_images(mock_post.call_args.kwargs["json"]))

    @patch("chat.model_manager.requests.Session.post", return_value=ok_response("coded"))
    def test_coding_question_routes_to_coder(self, mock_post):
        doc = Document.objects.create(original_name="r.pdf", file_type="PDF", extracted_text="Readings: 1, 2, 3")
        response = self.client.post(
            reverse("document-ask", kwargs={"pk": doc.id}),
            {"question": "Write a Python function that parses these readings."}, format="json",
        )
        self.assertEqual(response.data["ollama_model"], "qwen2.5-coder:7b")

    @patch("chat.model_manager.requests.Session.post", return_value=ok_response("1. Executive Overview"))
    def test_summary_uses_qwen3(self, mock_post):
        doc = Document.objects.create(original_name="s.pdf", file_type="PDF", extracted_text="Findings...")
        response = self.client.post(reverse("document-summary", kwargs={"pk": doc.id}), {}, format="json")
        self.assertEqual(response.data["ollama_model"], "qwen3:8b")
        self.assertEqual({k: v for k, v in response.data["routing"].items() if k != "explanation"}, {"mode": "fixed", "reason": "general_reasoning"})
        self.assertEqual(mock_post.call_args.kwargs["json"]["model"], "qwen3:8b")

    def test_ocr_module_has_no_model_execution_logic(self):
        import inspect
        from . import ocr
        source = inspect.getsource(ocr)
        for forbidden in ("import requests", "OLLAMA_URL", "/api/generate", "from chat", "import chat"):
            self.assertNotIn(forbidden, source)


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class DocumentLimitTests(APITestCase):
    """Phase A: configurable, reported limits (L16) and model-context fitting (L28)."""

    def upload(self, name, data, mime):
        return self.client.post(
            reverse("document-upload"),
            {"file": SimpleUploadedFile(name, data, content_type=mime)}, format="multipart",
        )

    @override_settings(DOCUMENT_MAX_CHARS=100)
    def test_char_limit_is_configurable_and_reported(self):
        response = self.upload("long.txt", ("abc " * 100).encode(), "text/plain")
        doc = response.data["document"]
        self.assertTrue(doc["truncated"])
        self.assertEqual(doc["original_chars"], 399)
        self.assertEqual(doc["limits"]["max_chars"], 100)

    @override_settings(DOCUMENT_MAX_PDF_PAGES=2)
    def test_pdf_page_limit_rejects_cleanly(self):
        doc = pymupdf.open()
        for _ in range(3):
            doc.new_page().insert_text((50, 72), "A typed page with enough characters to count as text.")
        response = self.upload("big.pdf", doc.tobytes(), "application/pdf")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("maximum is 2", response.data["error"])

    @override_settings(OCR_MAX_PDF_PAGES=1)
    @patch("documents.ocr.ocr_pdf_page", return_value=OCR_RESULT)
    def test_skipped_ocr_pages_reported_in_fields(self, mock_ocr):
        response = self.upload("scan.pdf", make_scanned_pdf_bytes(pages=3), "application/pdf")
        doc = response.data["document"]
        self.assertEqual((doc["page_count"], doc["ocr_pages"], doc["ocr_pages_skipped"]), (3, 1, 2))

    @patch("chat.model_manager.requests.Session.post", return_value=ok_response("ok"))
    def test_long_document_fitted_to_context_window(self, mock_post):
        text = "Reading 42 bar at valve. " * 2000  # 50,000 chars >> 8k-token window
        doc = Document.objects.create(original_name="long.txt", file_type="TXT", extracted_text=text)
        response = self.client.post(reverse("document-ask", kwargs={"pk": doc.id}),
                                    {"question": "What is the reading?"}, format="json")
        ctx = response.data["context"]
        self.assertTrue(ctx["context_truncated"])
        self.assertEqual(ctx["document_chars"], 50000)
        self.assertLess(ctx["context_chars_used"], 20000)
        prompt = chat_prompt(mock_post.call_args.kwargs["json"])
        self.assertLess(len(prompt) / 3.0, 8192)  # whole prompt fits, so Ollama will not halve it
        self.assertIn("What is the reading?", prompt)  # question survives
        self.assertIn("sections most relevant to the question", prompt)

    @patch("chat.model_manager.requests.Session.post", return_value=ok_response("ok"))
    def test_short_document_not_truncated(self, mock_post):
        doc = Document.objects.create(original_name="s.txt", file_type="TXT", extracted_text="Short text.")
        response = self.client.post(reverse("document-summary", kwargs={"pk": doc.id}), {}, format="json")
        self.assertFalse(response.data["context"]["context_truncated"])


class CleanupCommandTests(APITestCase):
    def test_dry_run_by_default_then_delete(self):
        from django.core.management import call_command, CommandError
        keep = Document.objects.create(original_name="real_report.pdf", file_type="PDF", extracted_text="x")
        test_doc = Document.objects.create(original_name="scanned_test.pdf", file_type="PDF", extracted_text="x")
        out = io.StringIO()
        call_command("cleanup_dev_data", "--name", "scanned_test", stdout=out)
        self.assertIn("Dry run", out.getvalue())
        self.assertEqual(Document.objects.count(), 2)
        call_command("cleanup_dev_data", "--name", "scanned_test", "--yes", stdout=io.StringIO())
        self.assertFalse(Document.objects.filter(pk=test_doc.pk).exists())
        self.assertTrue(Document.objects.filter(pk=keep.pk).exists())
        with self.assertRaises(CommandError):
            call_command("cleanup_dev_data", stdout=io.StringIO())  # no filter -> refuse


class DocumentConversationTests(APITestCase):
    """Phase C: Document Assistant Q&A keeps per-document history (L30)."""

    @patch("chat.model_manager.requests.Session.post")
    def test_document_follow_up_includes_previous_answer(self, mock_post):
        doc = Document.objects.create(original_name="r.txt", file_type="TXT",
                                      extracted_text="Finding: corrosion on E-2104A shell.")
        other = Document.objects.create(original_name="o.txt", file_type="TXT", extracted_text="Other.")
        url = reverse("document-ask", kwargs={"pk": doc.id})
        mock_post.return_value = ok_response("Corrosion on the E-2104A shell.")
        first = self.client.post(url, {"question": "What was found?"}, format="json")
        conv_id = first.data["conversation_id"]
        mock_post.return_value = ok_response("Within 30 days.")
        self.client.post(url, {"question": "How urgent is it?", "conversation_id": conv_id}, format="json")
        msgs = mock_post.call_args.kwargs["json"]["messages"]
        self.assertEqual([m["role"] for m in msgs], ["system", "user", "assistant", "user"])
        self.assertIn("corrosion on E-2104A shell", msgs[0]["content"])
        self.assertEqual(msgs[1]["content"], "What was found?")
        # A conversation belongs to one document: reusing its id on another starts fresh.
        self.client.post(reverse("document-ask", kwargs={"pk": other.id}),
                         {"question": "Anything?", "conversation_id": conv_id}, format="json")
        self.assertEqual(len(mock_post.call_args.kwargs["json"]["messages"]), 2)


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class OCRConfigurationTests(APITestCase):
    """Phase D: OCR profiles, preprocessing, languages, layout detection (L15, L18, L19, L21)."""

    @override_settings(OCR_MODEL_PROFILE="server")
    def test_default_engine_config(self):
        from .ocr import engine_config
        cfg = engine_config("en")
        self.assertEqual(cfg["lang"], "en")
        self.assertTrue(cfg["use_doc_orientation_classify"])   # fixes rotated scans
        self.assertFalse(cfg["use_textline_orientation"])
        self.assertFalse(cfg["enable_mkldnn"])
        self.assertEqual(cfg["device"], "cpu")

    @override_settings(OCR_MODEL_PROFILE="mobile")
    def test_mobile_profile(self):
        from .ocr import engine_config
        cfg = engine_config("en")
        self.assertEqual(cfg["text_detection_model_name"], "PP-OCRv5_mobile_det")
        self.assertEqual(cfg["text_recognition_model_name"], "en_PP-OCRv5_mobile_rec")
        self.assertEqual(engine_config("ka")["lang"], "ka")  # no mobile variant -> server/lang default

    @override_settings(OCR_DET_MODEL="X_det", OCR_REC_MODEL="X_rec", OCR_MODEL_PROFILE="mobile")
    def test_explicit_models_override_profile(self):
        from .ocr import engine_config
        cfg = engine_config("en")
        self.assertEqual((cfg["text_detection_model_name"], cfg["text_recognition_model_name"]), ("X_det", "X_rec"))

    @override_settings(OCR_ALLOWED_LANGS=["en", "ka"], OCR_LANG="en")
    @patch("documents.ocr.ocr_image", return_value=OCR_RESULT)
    def test_upload_language_choice(self, mock_ocr):
        url = reverse("document-upload")
        ok = self.client.post(url, {"file": SimpleUploadedFile("k.png", make_png_bytes(), content_type="image/png"),
                                    "ocr_lang": "ka"}, format="multipart")
        self.assertEqual(ok.status_code, status.HTTP_201_CREATED)
        self.assertEqual(ok.data["document"]["ocr_lang"], "ka")
        self.assertEqual(mock_ocr.call_args.args[1], "ka")
        bad = self.client.post(url, {"file": SimpleUploadedFile("k.png", make_png_bytes(), content_type="image/png"),
                                     "ocr_lang": "fr"}, format="multipart")
        self.assertEqual(bad.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("not enabled", bad.data["error"])
        cfg = self.client.get(reverse("document-ocr-config")).data
        self.assertEqual([l["code"] for l in cfg["languages"]], ["en", "ka"])
        self.assertEqual(cfg["default"], "en")

    @patch("documents.ocr.ocr_image", return_value={**OCR_RESULT, "complex_layout": True})
    def test_complex_layout_flag_stored(self, mock_ocr):
        response = self.client.post(reverse("document-upload"), {
            "file": SimpleUploadedFile("t.png", make_png_bytes(), content_type="image/png")}, format="multipart")
        self.assertTrue(response.data["document"]["complex_layout"])

    def test_layout_heuristic(self):
        from .ocr import detect_complex_layout
        single_column = [[100, 100 + i * 40, 900, 125 + i * 40] for i in range(10)]
        self.assertFalse(detect_complex_layout(single_column))
        table = []
        for i in range(6):
            y = 100 + i * 60
            table += [[100, y, 250, y + 25], [400, y, 500, y + 25], [700, y, 800, y + 25]]
        self.assertTrue(detect_complex_layout(table))
        one_split_line = single_column + [[100, 600, 300, 625], [600, 600, 900, 625]]
        self.assertFalse(detect_complex_layout(one_split_line))


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class DocumentVisionPagesTests(APITestCase):
    """Phase E: PDF visual analysis beyond page 1, never all pages at once (L17)."""

    def test_requested_pages(self):
        from .processing import requested_pages
        self.assertEqual(requested_pages("What does the drawing show?", 5), [1])
        self.assertEqual(requested_pages("Describe the diagram on page 3", 5), [3])
        self.assertEqual(requested_pages("Compare pages 2-4", 5), [2, 3])   # capped at 2
        self.assertEqual(requested_pages("pages 4 to 9", 5), [4, 5])
        self.assertEqual(requested_pages("page 12", 5), [1])               # out of range

    @patch("chat.model_manager.requests.Session.post", return_value=ok_response("A title block."))
    def test_page_reference_selects_rendered_page(self, mock_post):
        doc = Document.objects.create(
            original_name="scan.pdf", file_type="PDF", extraction_method="ocr", extracted_text="x" * 30,
            page_count=3, file=SimpleUploadedFile("scan.pdf", make_scanned_pdf_bytes(pages=3)),
        )
        response = self.client.post(reverse("document-ask", kwargs={"pk": doc.id}),
                                    {"question": "What does the drawing on page 3 show?"}, format="json")
        self.assertEqual(response.data["vision_pages"], [3])
        self.assertEqual(response.data["page_count"], 3)
        payload = mock_post.call_args.kwargs["json"]
        self.assertEqual(len(chat_images(payload)), 1)
        self.assertIn("page 3 of 3", payload["messages"][0]["content"])
        both = self.client.post(reverse("document-ask", kwargs={"pk": doc.id}),
                                {"question": "Describe the drawings on pages 1-3"}, format="json")
        self.assertEqual(both.data["vision_pages"], [1, 2])
        self.assertEqual(len(chat_images(mock_post.call_args.kwargs["json"])), 2)


class DocumentStreamRegenerateTests(APITestCase):
    """Limitations follow-up: streaming + regenerate for Document Q&A."""

    def setUp(self):
        self.doc = Document.objects.create(original_name="r.txt", file_type="TXT",
                                           extracted_text="Finding: shell corrosion on E-2104A.")

    @patch("chat.model_manager.requests.Session.post")
    def test_stream_answer(self, mock_post):
        import json as _json
        from chat.tests import ollama_stream
        mock_post.return_value = ollama_stream("Shell ", "corrosion.")
        r = self.client.post(reverse("document-ask-stream", kwargs={"pk": self.doc.id}),
                             {"question": "What was found?"}, format="json")
        events = [_json.loads(l) for l in b"".join(r.streaming_content).decode().splitlines() if l]
        self.assertEqual(events[0]["type"], "start")
        self.assertEqual(events[0]["routing"]["reason"], "general_reasoning")
        self.assertEqual(events[-1]["type"], "done")
        self.assertEqual(events[-1]["answer"], "Shell corrosion.")
        self.assertEqual(events[-1]["ollama_model"], "qwen3:8b")

    @patch("chat.model_manager.requests.Session.post")
    def test_regenerate_document_answer(self, mock_post):
        url = reverse("document-ask", kwargs={"pk": self.doc.id})
        mock_post.return_value = ok_response("first")
        conv_id = self.client.post(url, {"question": "What was found?"}, format="json").data["conversation_id"]
        mock_post.return_value = ok_response("second")
        r = self.client.post(url, {"regenerate": True, "conversation_id": conv_id}, format="json")
        self.assertEqual(r.data["answer"], "second")
        self.assertEqual([m["role"] for m in mock_post.call_args.kwargs["json"]["messages"]], ["system", "user"])
        msgs = self.client.get(reverse("conversation-detail", kwargs={"pk": conv_id})).data["conversation"]["messages"]
        self.assertEqual([m["content"] for m in msgs], ["What was found?", "second"])
        self.assertEqual(self.client.post(url, {"regenerate": True}, format="json").status_code, 400)


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class DocumentReliabilityTests(APITestCase):
    @patch("chat.model_manager.requests.Session.post")
    def test_document_grounding_prompt_and_basis(self, mock_post):
        doc = Document.objects.create(original_name="plant.txt", file_type="TXT",
                                      extracted_text="Project location: Jaipur\nMaximum operating temperature: 120°C")
        mock_post.return_value = ok_response("I could not find this information in the provided document.")
        r = self.client.post(reverse("document-ask", kwargs={"pk": doc.id}),
                             {"question": "What is the project budget?"}, format="json")
        system = mock_post.call_args.kwargs["json"]["messages"][0]["content"]
        self.assertIn("Answer only from the supplied document content", system)
        self.assertIn("Maximum operating temperature: 120°C", system)
        self.assertEqual(r.data["basis"]["label"], "Not found in the provided document")
        mock_post.return_value = ok_response("The maximum operating temperature is 120°C.")
        r = self.client.post(reverse("document-ask", kwargs={"pk": doc.id}),
                             {"question": "What is the maximum operating temperature?"}, format="json")
        self.assertEqual((r.data["basis"]["label"], r.data["basis"]["detail"]), ("Based on uploaded document", "plant.txt"))

    @patch("chat.model_manager.requests.Session.post", return_value=ok_response("A title block."))
    def test_visual_basis_names_the_pages_actually_sent(self, mock_post):
        doc = Document.objects.create(
            original_name="scan.pdf", file_type="PDF", extraction_method="ocr", extracted_text="x" * 30,
            page_count=3, file=SimpleUploadedFile("scan.pdf", make_scanned_pdf_bytes(pages=3)),
        )
        r = self.client.post(reverse("document-ask", kwargs={"pk": doc.id}),
                             {"question": "What does the drawing on page 2 show?"}, format="json")
        self.assertEqual(r.data["basis"]["detail"], "scan.pdf — page 2 (image)")
        self.assertIn("Observed", mock_post.call_args.kwargs["json"]["messages"][0]["content"])


_sync_override = override_settings(DOCUMENT_ASYNC_PROCESSING=False)


def setUpModule():
    from chat.tests import start_loaded_patchers
    start_loaded_patchers()
    _sync_override.enable()  # most tests check the finished upload; async is tested separately


def tearDownModule():
    from chat.tests import stop_loaded_patchers
    stop_loaded_patchers()
    _sync_override.disable()


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class AsyncUploadTests(APITransactionTestCase):
    """The worker thread writes through its own DB connection, so no wrapping transaction."""

    @override_settings(DOCUMENT_ASYNC_PROCESSING=True)
    @patch("documents.ocr.ocr_pdf_page", return_value=OCR_RESULT)
    def test_async_upload_then_ready(self, mock_ocr):
        import time as _time
        from . import jobs
        r = self.client.post(reverse("document-upload"), {
            "file": SimpleUploadedFile("scan.pdf", make_scanned_pdf_bytes(pages=2), content_type="application/pdf")},
            format="multipart")
        self.assertEqual(r.status_code, 202)
        doc_id = r.data["document"]["id"]
        self.assertIn(r.data["document"]["status"], ("queued", "processing", "ready"))
        for _ in range(100):  # the worker thread runs in the background
            if doc_id not in jobs.active_ids():
                break
            _time.sleep(0.05)
        d = self.client.get(reverse("document-detail", kwargs={"pk": doc_id})).data["document"]
        self.assertEqual(d["status"], "ready")
        self.assertEqual((d["progress_done"], d["progress_total"]), (2, 2))
        self.assertIn("shell corrosion", d["extracted_text"])


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class BackgroundProcessingTests(APITestCase):
    """Response time: uploads return immediately, extraction runs in a worker (L31)."""

    def test_questions_blocked_until_ready_and_interrupted_jobs_fail(self):
        doc = Document.objects.create(original_name="p.pdf", file_type="PDF", extracted_text="",
                                      status=Document.STATUS_PROCESSING)
        from . import jobs
        with patch.object(jobs, "active_ids", return_value={doc.id}):
            r = self.client.post(reverse("document-ask", kwargs={"pk": doc.id}), {"question": "?"}, format="json")
            self.assertEqual(r.status_code, 409)
            self.assertIn("still being processed", r.data["error"])
        # No job exists for it any more (e.g. server restarted) -> reported as failed, not stuck.
        d = self.client.get(reverse("document-detail", kwargs={"pk": doc.id})).data["document"]
        self.assertEqual(d["status"], "failed")
        self.assertIn("interrupted", d["processing_error"])

    def test_page_offsets_for_text_and_ocr_pdfs(self):
        from .extractors import extract_document
        src = pymupdf.open()
        for i in range(3):
            src.new_page().insert_text((50, 72), f"Typed page number {i + 1} with enough characters here.")
        res = extract_document(io.BytesIO(src.tobytes()), "PDF")
        for page, start in res["page_offsets"]:
            self.assertTrue(res["text"][start:].startswith(f"Typed page number {page}"))
        with patch("documents.ocr.ocr_pdf_page", return_value=OCR_RESULT):
            res = extract_document(io.BytesIO(make_scanned_pdf_bytes(pages=2)), "PDF")
        self.assertEqual([p for p, _ in res["page_offsets"]], [1, 2])
        self.assertTrue(res["text"][res["page_offsets"][1][1]:].startswith("--- Page 2 (OCR) ---"))

    def test_generic_error_on_unexpected_failure(self):
        with patch("documents.extractors.extract_document", side_effect=RuntimeError("secret /var/path")):
            r = self.client.post(reverse("document-upload"), {
                "file": SimpleUploadedFile("a.txt", b"hello world", content_type="text/plain")}, format="multipart")
        self.assertEqual(r.status_code, 400)
        self.assertNotIn("secret", r.data["error"])


class RetrievalAndVerificationTests(APITestCase):
    """L28 (long documents), page citations, deterministic number/citation checks."""

    def long_pdf_doc(self):
        from .extractors import extract_document
        src = pymupdf.open()
        for i in range(1, 41):
            page = src.new_page()
            body = "Routine inspection log entry. " * 40
            if i == 33:
                body += " The emergency shutdown pressure for vessel V-101 is 17.5 bar."
            page.insert_textbox(pymupdf.Rect(40, 40, 560, 800), f"Page {i} report.\n{body}", fontsize=9)
        res = extract_document(io.BytesIO(src.tobytes()), "PDF")
        return Document.objects.create(original_name="long.pdf", file_type="PDF", extracted_text=res["text"],
                                       page_offsets=res["page_offsets"], page_count=40)

    @patch("chat.model_manager.requests.Session.post")
    def test_fact_deep_in_long_document_is_retrieved_with_page_label(self, mock_post):
        doc = self.long_pdf_doc()
        self.assertGreater(len(doc.extracted_text), 40000)
        mock_post.return_value = ok_response("The shutdown pressure is 17.5 bar [p. 33].")
        r = self.client.post(reverse("document-ask", kwargs={"pk": doc.id}),
                             {"question": "What is the emergency shutdown pressure for vessel V-101?"}, format="json")
        system = mock_post.call_args.kwargs["json"]["messages"][0]["content"]
        self.assertIn("[p. 33]", system)
        self.assertIn("17.5 bar", system)
        self.assertTrue(r.data["context"]["retrieval"])
        self.assertIn(33, r.data["context"]["pages"])
        self.assertEqual(r.data["basis"]["detail"], "long.pdf — p. 33")
        self.assertEqual(r.data["basis"]["warnings"], [])

    @patch("chat.model_manager.requests.Session.post")
    def test_invented_value_and_bad_citation_are_flagged(self, mock_post):
        doc = self.long_pdf_doc()
        mock_post.return_value = ok_response("The shutdown pressure is 23.75 bar [p. 39].")
        r = self.client.post(reverse("document-ask", kwargs={"pk": doc.id}),
                             {"question": "What is the emergency shutdown pressure for vessel V-101?"}, format="json")
        w = " ".join(r.data["basis"]["warnings"])
        self.assertIn("not found in the document: 23.75", w)
        self.assertIn("Cited page not given to the model: 39", w)
        stored = self.client.get(reverse("conversation-detail", kwargs={"pk": r.data["conversation_id"]}))
        self.assertEqual(stored.data["conversation"]["messages"][1]["basis"]["warnings"], r.data["basis"]["warnings"])

    def test_number_check_rules(self):
        from .verify import check_numbers
        doc = "Design pressure: 16 bar(g). Completion Q3 2027. Temperature 120°C. Flow 1,250 m3/h."
        self.assertEqual(check_numbers("It is 16 bar and 120°C, flow 1250 m3/h.", doc), [])
        self.assertEqual(check_numbers("1. Pressure 16 bar\n2. Budget 5 crore", doc), ["5"])
        self.assertEqual(check_numbers("Page cited [p. 4]: 16 bar", doc), [])

    def test_summary_spreads_over_whole_document(self):
        from .retrieval import build_context
        text = "\n\n".join(f"Section {i}: " + "text " * 200 for i in range(30))
        content, info = build_context(text, "", 6000)
        self.assertTrue(info["retrieval"])
        self.assertIn("Section 0:", content)
        self.assertIn("Section 2", content)
        self.assertTrue(any(f"Section {i}:" in content for i in range(20, 30)))  # not only the beginning

    def test_short_document_passed_whole(self):
        from .retrieval import build_context
        content, info = build_context("Short text.", "question", 5000)
        self.assertEqual((content, info["retrieval"]), ("Short text.", False))


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class LayoutFallbackOverviewTests(APITestCase):
    """L19 tables/columns, L20/L18 vision fallback, L17 whole-document visual overview."""

    def test_table_reconstruction(self):
        from .ocr import reconstruct_layout
        texts, boxes = ["Readings"], [[100, 20, 400, 45]]
        for r, row in enumerate([["Location", "Nominal", "Measured"], ["North", "10.0", "9.6"], ["South", "10.0", "7.2"]]):
            y = 100 + r * 60
            for c, cell in enumerate(row):
                texts.append(cell)
                boxes.append([100 + c * 300, y, 200 + c * 300, y + 25])
        text, kind = reconstruct_layout(texts, boxes)
        self.assertEqual(kind, "table")
        self.assertIn("| Location | Nominal | Measured |", text)
        self.assertIn("| South | 10.0 | 7.2 |", text)

    def test_two_columns_read_in_order(self):
        from .ocr import reconstruct_layout
        texts, boxes = [], []
        for i in range(5):
            texts += [f"L{i}", f"R{i}"]
            boxes += [[80, 100 + i * 40, 500, 125 + i * 40], [680, 100 + i * 40, 1100, 125 + i * 40]]
        text, kind = reconstruct_layout(texts, boxes)
        self.assertEqual(kind, "columns")
        self.assertEqual(text.split(), ["L0", "L1", "L2", "L3", "L4", "R0", "R1", "R2", "R3", "R4"])

    @patch("chat.model_manager.requests.Session.post")
    @patch("documents.ocr.ocr_image", return_value={"text": "Pnmp P-22l0B vbratn", "lines": 1,
                                                   "confidence": 0.61, "complex_layout": False})
    def test_low_confidence_ocr_gets_vision_transcription(self, mock_ocr, mock_post):
        mock_post.return_value = ok_response("Pump P-2210B vibration high.")
        png = make_png_bytes(400, 200)
        r = self.client.post(reverse("document-upload"), {
            "file": SimpleUploadedFile("note.png", png, content_type="image/png")}, format="multipart")
        doc = r.data["document"]
        self.assertEqual(doc["vision_fallback_pages"], 1)
        self.assertIn("(vision transcription)", doc["extracted_text"])
        self.assertIn("Pump P-2210B vibration high.", doc["extracted_text"])
        payload = mock_post.call_args.kwargs["json"]
        self.assertEqual(payload["model"], "qwen2.5vl:7b")
        self.assertIn("Transcribe all readable text", payload["prompt"])

    @override_settings(OCR_VISION_FALLBACK=False)
    @patch("documents.ocr.ocr_image", return_value={"text": "blurry", "lines": 1, "confidence": 0.5,
                                                   "complex_layout": False})
    def test_fallback_can_be_disabled(self, mock_ocr):
        r = self.client.post(reverse("document-upload"), {
            "file": SimpleUploadedFile("note.png", make_png_bytes(), content_type="image/png")}, format="multipart")
        self.assertEqual(r.data["document"]["vision_fallback_pages"], 0)

    def test_overview_detection_and_batches(self):
        from .processing import wants_overview, overview_batches
        self.assertTrue(wants_overview("Describe the drawings on all pages"))
        self.assertTrue(wants_overview("Summarise the whole document page by page"))
        self.assertFalse(wants_overview("What does page 2 show?"))
        self.assertEqual(overview_batches(9), ([[1, 2], [3, 4], [5, 6]], 3))
        self.assertEqual(overview_batches(3), ([[1, 2], [3]], 0))

    def _pdf_doc(self, pages=5):
        return Document.objects.create(
            original_name="set.pdf", file_type="PDF", extraction_method="ocr", extracted_text="x" * 30,
            page_count=pages, file=SimpleUploadedFile("set.pdf", make_scanned_pdf_bytes(pages=pages)),
        )

    @patch("chat.model_manager.requests.Session.post")
    def test_overview_json_runs_batches_sequentially(self, mock_post):
        doc = self._pdf_doc(5)
        mock_post.return_value = ok_response("A blank page.")
        r = self.client.post(reverse("document-ask", kwargs={"pk": doc.id}),
                             {"question": "Describe the drawings on all pages."}, format="json")
        self.assertEqual(mock_post.call_count, 3)  # pages 1-2, 3-4, 5
        self.assertEqual([len(c.kwargs["json"]["messages"][-1]["images"]) for c in mock_post.call_args_list], [2, 2, 1])
        self.assertIn("### Pages 1–2", r.data["answer"])
        self.assertIn("### Page 5", r.data["answer"])
        self.assertEqual(r.data["vision_pages"], [1, 2, 3, 4, 5])
        self.assertIn("page by page", r.data["basis"]["detail"])

    @patch("chat.model_manager.requests.Session.post")
    def test_overview_stream(self, mock_post):
        import json as _json
        from chat.tests import ollama_stream
        doc = self._pdf_doc(8)
        mock_post.side_effect = lambda *a, **k: ollama_stream("Blank.")
        r = self.client.post(reverse("document-ask-stream", kwargs={"pk": doc.id}),
                             {"question": "Describe the drawings on every page."}, format="json")
        events = [_json.loads(l) for l in b"".join(r.streaming_content).decode().splitlines() if l]
        self.assertEqual(events[0]["pages_not_covered"], 2)  # 8 pages, limit 6
        done = events[-1]
        self.assertEqual(done["type"], "done")
        self.assertEqual(done["answer"].count("###"), 3)
